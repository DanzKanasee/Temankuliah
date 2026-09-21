import os
from pypdf import PdfReader
from docx import Document
from pptx import Presentation

def extract_text_from_file(file_path: str) -> str:
    """Mengekstrak teks dari file PDF, DOCX, atau PPTX berdasarkan ekstensi."""
    ext = os.path.splitext(file_path)[1].lower()
    text = ""

    try:
        if ext == ".pdf":
            reader = PdfReader(file_path)
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"

        elif ext in [".docx", ".doc"]:
            doc = Document(file_path)
            for paragraph in doc.paragraphs:
                if paragraph.text:
                    text += paragraph.text + "\n"

        elif ext in [".pptx", ".ppt"]:
            prs = Presentation(file_path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        text += shape.text + "\n"

        return text.strip()
    except Exception as e:
        print(f"Error extracting text from {file_path}: {e}")
        return ""
